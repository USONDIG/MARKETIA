from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from business_themes import add_business_theme_columns
from database import connect
from discovery import build_discovery
from qualification import employee_info, naf_division

OUTPUT_DIR = Path("output")
XLSX_PATH = OUTPUT_DIR / "server_infra_radar.xlsx"
ALERTS_PATH = OUTPUT_DIR / "alerts.csv"


def _query_df(sql: str) -> pd.DataFrame:
    with connect() as db:
        return pd.read_sql_query(sql, db)


def _add_qualification_columns(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    df = df.copy()
    df["naf_division"] = df["naf"].apply(naf_division)
    df["employee_min"] = df["employees"].apply(lambda value: employee_info(value)["employee_min"])
    df["employee_range"] = df["employees"].apply(lambda value: employee_info(value)["employee_range"])
    df["qualified_target"] = True
    return df


def export_outputs(config: dict) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    threshold = float(config.get("alerts", {}).get("minimum_score", 65))
    top_n = int(config.get("automation", {}).get("excel_top_leads", 500))
    discovery_n = int(config.get("automation", {}).get("excel_top_discovery", 1000))

    opportunities = _query_df(
        f"""
        SELECT
            s.priority,
            s.opportunity_score,
            s.infra_fit,
            s.buying_intent,
            s.timing,
            s.company_name,
            s.siren,
            c.naf,
            c.city,
            c.employees,
            c.sites,
            c.headquarters_country,
            s.top_signal,
            s.calculated_at
        FROM scores s
        LEFT JOIN companies c ON c.siren = s.siren
        ORDER BY s.opportunity_score DESC
        LIMIT {top_n}
        """
    )
    opportunities = _add_qualification_columns(opportunities)
    opportunities = add_business_theme_columns(opportunities)
    discovery = build_discovery(config, limit=discovery_n)

    events = _query_df(
        """
        SELECT
            e.event_date, e.source, e.event_type, e.company_name, e.siren,
            e.country, e.title, e.url, e.collected_at,
            GROUP_CONCAT(DISTINCT s.label) AS detected_signals
        FROM events e
        LEFT JOIN signals s ON s.event_id = e.id
        GROUP BY e.id
        ORDER BY COALESCE(e.event_date, e.collected_at) DESC
        LIMIT 5000
        """
    )

    signals = _query_df(
        """
        SELECT
            e.company_name, e.siren, e.source, e.event_date,
            s.signal_type, s.label, s.strength, s.weight,
            s.matched_terms, e.title
        FROM signals s
        JOIN events e ON e.id = s.event_id
        ORDER BY s.detected_at DESC
        LIMIT 10000
        """
    )

    runs = _query_df("SELECT * FROM runs ORDER BY source")
    alerts = opportunities[opportunities["opportunity_score"] >= threshold].copy() if not opportunities.empty else opportunities.copy()
    alerts.to_csv(ALERTS_PATH, index=False, encoding="utf-8-sig")

    discovery_qualified = int((discovery["status"] == "QUALIFIED").sum()) if not discovery.empty else 0
    dashboard = pd.DataFrame([
        ["Radar / Discovery", int(len(discovery))],
        ["Discovery qualifies", discovery_qualified],
        ["Leads qualifies", int(len(opportunities))],
        ["CRITICAL", int((opportunities["priority"] == "CRITICAL").sum()) if not opportunities.empty else 0],
        ["HOT", int((opportunities["priority"] == "HOT").sum()) if not opportunities.empty else 0],
        ["WARM", int((opportunities["priority"] == "WARM").sum()) if not opportunities.empty else 0],
        ["Alertes >= seuil", int(len(alerts))],
        ["Cible qualifiee", "Industrie + commerce + services + administration publique, FR, 50+ salaries"],
        ["Seuil d'alerte", threshold],
    ], columns=["Indicateur", "Valeur"])

    with pd.ExcelWriter(XLSX_PATH, engine="openpyxl") as writer:
        dashboard.to_excel(writer, sheet_name="Dashboard", index=False, startrow=2)
        discovery.to_excel(writer, sheet_name="Discovery", index=False)
        opportunities.to_excel(writer, sheet_name="Opportunities", index=False)
        alerts.to_excel(writer, sheet_name="Alerts", index=False)
        events.to_excel(writer, sheet_name="Events", index=False)
        signals.to_excel(writer, sheet_name="Signals", index=False)
        runs.to_excel(writer, sheet_name="Runs", index=False)

    wb = load_workbook(XLSX_PATH)
    dark = PatternFill("solid", fgColor="111827")
    header = PatternFill("solid", fgColor="374151")
    white_font = Font(color="FFFFFF", bold=True)

    ws = wb["Dashboard"]
    ws["A1"] = "MARKETIA — SERVER INFRA OPPORTUNITY RADAR"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = dark
    ws.merge_cells("A1:D1")
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2" if sheet.title != "Dashboard" else "A4"
        header_row = 1 if sheet.title != "Dashboard" else 3
        for cell in sheet[header_row]:
            cell.fill = header
            cell.font = white_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for col_idx in range(1, sheet.max_column + 1):
            max_len = 0
            for row_idx in range(1, min(sheet.max_row, 200) + 1):
                value = sheet.cell(row_idx, col_idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 42)

    if "Opportunities" in wb.sheetnames and wb["Opportunities"].max_row > 1:
        rng = f"B2:B{wb['Opportunities'].max_row}"
        wb["Opportunities"].conditional_formatting.add(
            rng,
            ColorScaleRule(start_type="num", start_value=0, start_color="FEE2E2",
                           mid_type="num", mid_value=65, mid_color="FEF3C7",
                           end_type="num", end_value=100, end_color="DCFCE7")
        )

    if "Discovery" in wb.sheetnames and wb["Discovery"].max_row > 1:
        rng = f"B2:B{wb['Discovery'].max_row}"
        wb["Discovery"].conditional_formatting.add(
            rng,
            ColorScaleRule(start_type="num", start_value=0, start_color="FEE2E2",
                           mid_type="num", mid_value=60, mid_color="FEF3C7",
                           end_type="num", end_value=100, end_color="DCFCE7")
        )

    wb.save(XLSX_PATH)
